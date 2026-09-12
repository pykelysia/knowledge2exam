"""DOCX 解析器：使用 python-docx 提取文本和内嵌图片。"""

from __future__ import annotations

import io

from app.ingestion.parsers.base import ImageInfo, Parser, ParseResult


class DocxParser(Parser):
    """DOCX 解析器，提取段落文本和内嵌图片。"""

    def parse(self, filename: str, data: bytes) -> ParseResult:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("python-docx 未安装，请运行 `pip install python-docx`") from exc

        doc = Document(io.BytesIO(data))

        # 提取所有段落文本
        parts: list[str] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)

        # 也提取表格内容
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)

        # 提取内嵌图片（通过 document.part 的关系）
        images: list[ImageInfo] = []
        seen_rels: set[str] = set()
        try:
            for rel in doc.part.rels.values():
                if "image" not in rel.reltype:
                    continue
                if rel.rId in seen_rels:
                    continue
                seen_rels.add(rel.rId)
                try:
                    image_part = rel.target_part
                    image_bytes = image_part.blob
                    images.append(
                        ImageInfo(
                            page=None,
                            data=image_bytes,
                            order=len(images),
                        )
                    )
                except Exception:
                    continue
        except Exception:
            pass

        text = "\n\n".join(parts)

        return ParseResult(
            char_count=len(text),
            page_count=None,
            text=text,
            images=images if images else None,
        )
