"""PPTX 解析器：python-pptx 原生提取（标题/文本/表格/图片占位符/备注）。

输出统一文本契约：幻灯片标题 → `#` 一级标题；页码/日期/页脚占位符跳过；
备注存在则作为该页附加段落；图片在正文位置留 `[图: slideN-M]` 占位符
（供编排层图片 OCR 原位回填）。OCR 前过滤装饰图：跨页重复（logo/水印）
与整页背景图。
"""

from __future__ import annotations

import hashlib
import io

from app.ingestion.parsers.base import ImageInfo, Parser, ParseResult
from app.ingestion.parsers.tables import rows_to_markdown

# 同字节图片跨页出现次数达到该值视为 logo/水印
_LOGO_REPEAT_THRESHOLD = 2
# 图片宽高均达到幻灯片尺寸的该比例视为背景图
_BACKGROUND_COVERAGE = 0.95


class PPTXParser(Parser):
    """PPTX 解析器，逐页提取标题、正文、表格、图片与备注。"""

    def parse(self, filename: str, data: bytes) -> ParseResult:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError("python-pptx 未安装，请运行 `pip install python-pptx`") from exc

        prs = Presentation(io.BytesIO(data))
        slides = list(prs.slides)

        # 第一遍：统计图片字节指纹，识别跨页重复的 logo/水印
        blob_counts: dict[str, int] = {}
        for slide in slides:
            for shape in _iter_shapes(slide.shapes):
                if _is_picture(shape):
                    try:
                        blob = shape.image.blob
                    except Exception:  # noqa: BLE001
                        continue
                    digest = hashlib.sha256(blob).hexdigest()
                    blob_counts[digest] = blob_counts.get(digest, 0) + 1

        blocks: list[str] = []
        images: list[ImageInfo] = []
        for slide_num, slide in enumerate(slides, start=1):
            blocks.extend(self._parse_slide(slide, slide_num, prs, blob_counts, images))

        text = "\n\n".join(blocks)

        return ParseResult(
            char_count=len(text),
            page_count=len(slides),
            text=text,
            images=images or None,
        )

    def _parse_slide(
        self,
        slide,
        slide_num: int,
        prs,
        blob_counts: dict[str, int],
        images: list[ImageInfo],
    ) -> list[str]:
        """解析单页，返回该页的块序列（标题块在最前，备注块在最后）。"""
        title: str | None = None
        body: list[str] = []

        for shape in _iter_shapes(slide.shapes):
            if _is_picture(shape):
                self._collect_picture(shape, slide_num, prs, blob_counts, images, body)
                continue
            if getattr(shape, "has_table", False):
                markdown = rows_to_markdown(
                    [[cell.text for cell in row.cells] for row in shape.table.rows]
                )
                if markdown:
                    body.append(markdown)
                continue
            if not shape.has_text_frame:
                continue
            kind = _placeholder_kind(shape)
            if kind == "title":
                text = shape.text_frame.text.strip()
                if text:
                    title = text
                continue
            if kind == "ignore":
                continue
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    body.append(text)

        blocks: list[str] = []
        if title:
            blocks.append(f"# {title}")
        blocks.extend(body)
        notes = _slide_notes(slide)
        if notes:
            blocks.append(f"备注：{notes}")
        return blocks

    @staticmethod
    def _collect_picture(
        shape,
        slide_num: int,
        prs,
        blob_counts: dict[str, int],
        images: list[ImageInfo],
        body: list[str],
    ) -> None:
        """装饰图（logo/背景）跳过；内容图登记占位符并就地输出。"""
        try:
            blob = shape.image.blob
        except Exception:  # noqa: BLE001
            return
        digest = hashlib.sha256(blob).hexdigest()
        if blob_counts.get(digest, 0) >= _LOGO_REPEAT_THRESHOLD:
            return
        if (
            shape.width
            and shape.height
            and shape.width >= prs.slide_width * _BACKGROUND_COVERAGE
            and shape.height >= prs.slide_height * _BACKGROUND_COVERAGE
        ):
            return
        marker = f"[图: slide{slide_num}-{len(images) + 1}]"
        images.append(ImageInfo(page=slide_num, data=blob, order=len(images) + 1, marker=marker))
        body.append(marker)


def _iter_shapes(shapes):
    """展平组合形状，按文档顺序产出叶子 shape。"""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _is_picture(shape) -> bool:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    return shape.shape_type == MSO_SHAPE_TYPE.PICTURE


def _placeholder_kind(shape) -> str:
    """占位符语义分流：title / ignore（页码/日期/页脚）/ body。"""
    if not shape.is_placeholder:
        return "body"
    from pptx.enum.shapes import PP_PLACEHOLDER

    kind = shape.placeholder_format.type
    if kind in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
        return "title"
    if kind in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER):
        return "ignore"
    return "body"


def _slide_notes(slide) -> str:
    """演讲者备注；无备注页或读取失败返回空串。"""
    try:
        if not slide.has_notes_slide:
            return ""
        return (slide.notes_slide.notes_text_frame.text or "").strip()
    except Exception:  # noqa: BLE001
        return ""
